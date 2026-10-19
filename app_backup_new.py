# -*- coding: utf-8 -*-
from __future__ import annotations
import hashlib
import base64
import json
import urllib.request
import urllib.error
from datetime import datetime, date
from pathlib import Path
import pandas as pd
import streamlit as st

BASE = Path(__file__).parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)

FILES = {
    "employees": ("الموظفون", "employees.xlsx"),
    "contracts": ("عقود الموظفين", "contracts.xlsx"),
    "iqama": ("الإقامات", "iqama.xlsx"),
    "baladi": ("الكروت البلدية", "baladi_cards.xlsx"),
    "rents": ("الإيجارات", "rents.xlsx"),
    "movements": ("سجل الحركات", "movements.xlsx"),
    "archive": ("أرشيف الموظفين", "archive.xlsx"),
    "changes": ("سجل التعديلات", "changes.xlsx"),
}

ALERT_DAYS = {"contracts": 90, "iqama": 90, "passport": 90, "baladi": 20, "rents": 30}

st.set_page_config(page_title="نظام الموارد البشرية", page_icon="👥", layout="wide")
st.markdown("""
<style>
html,body,[class*="css"]{direction:rtl;text-align:right}
.stApp{background:linear-gradient(180deg,#06101d,#0a1422);color:#f8fafc}
[data-testid="stSidebar"]{background:#07101c;border-left:1px solid #243449}
[data-testid="stSidebar"] *{color:#e5edf7}
.block-container{padding-top:1rem;padding-bottom:2rem;max-width:1600px}
h1,h2,h3,p,label{color:#f8fafc!important}
[data-testid="stMetric"]{background:linear-gradient(145deg,#101c2d,#0d1725);border:1px solid #243449;border-radius:18px;padding:18px;box-shadow:0 10px 25px rgba(0,0,0,.22)}
[data-testid="stMetricLabel"] p{color:#cbd5e1!important}
[data-testid="stMetricValue"]{color:white}
.hero{background:linear-gradient(120deg,#111f33,#0b1727);border:1px solid #26384f;border-radius:22px;padding:22px 25px;margin-bottom:16px;box-shadow:0 14px 32px rgba(0,0,0,.25)}
.empty{background:#0e1928;border:1px dashed #334155;border-radius:16px;color:#94a3b8;padding:30px;text-align:center}
div[data-baseweb="input"]>div,div[data-baseweb="select"]>div{background:#101c2d!important;color:white!important;border-color:#334155!important}
color:black!important;caret-color:black!importantcolor:white!important}
.stButton button,.stDownloadButton button{border-radius:12px;font-weight:700;border:1px solid #334155}
[data-testid="stDataFrame"]{border:1px solid #26384f;border-radius:14px;overflow:hidden}
</style>
""", unsafe_allow_html=True)

def hpwd(v): return hashlib.sha256(v.encode()).hexdigest()
USERS = {
    "hr": {"password": hpwd("1234"), "role": "hr", "name": "HR"},
    "manager": {"password": hpwd("Manager@2026"), "role": "manager", "name": "Manager"},
}

def load_df(key):
    """Read authoritative saved data, not stale XLSX files from a deployment."""
    p = DATA / FILES[key][1]
    status, payload = _github_read_file(f"data/{FILES[key][1]}")
    if status == "found":
        try:
            from io import BytesIO
            return pd.read_excel(BytesIO(payload))
        except Exception as e:
            st.error(f"تعذر فتح النسخة المحفوظة من {FILES[key][0]}: {e}")
            return pd.DataFrame()
    if status == "missing":
        # A removed file in GitHub must never reappear from bundled app files.
        return pd.DataFrame()
    # No remote access: don't mistake a bundled old file for current data.
    st.warning(f"لا يمكن التحقق من النسخة الحالية لملف {FILES[key][0]}: {payload}")
    return pd.DataFrame()

GITHUB_OWNER = "be22-2"
GITHUB_REPO = "pastahouse-hr-rentals"
GITHUB_BRANCH = "main"
# GitHub is the source of truth for EVERY workbook (including deletions).
GITHUB_PERSIST_KEYS = set(FILES)

def _github_token():
    try:
        return str(st.secrets.get("github_token", "")).strip()
    except Exception:
        return ""

def _github_put_file(repo_path, local_path, message):
    """Upload/replace one file in GitHub using a token stored in Streamlit Secrets."""
    token = _github_token()
    if not token:
        return False, "GitHub token غير موجود في Streamlit Secrets"

    api_url = f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPO}/contents/{repo_path}"
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "pastahouse-hr-rentals",
    }

    sha = None
    try:
        req = urllib.request.Request(f"{api_url}?ref={GITHUB_BRANCH}", headers=headers, method="GET")
        with urllib.request.urlopen(req, timeout=20) as r:
            sha = json.loads(r.read().decode("utf-8")).get("sha")
    except urllib.error.HTTPError as e:
        if e.code != 404:
            try:
                detail = e.read().decode("utf-8", errors="ignore")
            except Exception:
                detail = str(e)
            return False, f"GitHub GET failed ({e.code}): {detail[:200]}"
    except Exception as e:
        return False, f"GitHub GET failed: {e}"

    payload = {
        "message": message,
        "content": base64.b64encode(Path(local_path).read_bytes()).decode("ascii"),
        "branch": GITHUB_BRANCH,
    }
    if sha:
        payload["sha"] = sha

    data = json.dumps(payload).encode("utf-8")
    try:
        req = urllib.request.Request(api_url, data=data, headers={**headers, "Content-Type": "application/json"}, method="PUT")
        with urllib.request.urlopen(req, timeout=30) as r:
            if r.status not in (200, 201):
                return False, f"GitHub PUT returned {r.status}"
        return True, ""
    except urllib.error.HTTPError as e:
        try:
            detail = e.read().decode("utf-8", errors="ignore")
        except Exception:
            detail = str(e)
        return False, f"GitHub PUT failed ({e.code}): {detail[:250]}"
    except Exception as e:
        return False, f"GitHub PUT failed: {e}"

@st.cache_data(ttl=30, show_spinner=False)
def _github_read_file(repo_path):
    """Return (found, bytes), (missing, None), or (error, explanation)."""
    token = _github_token()
    if not token:
        return "error", "مفتاح github_token غير موجود في Streamlit Secrets"
    url = f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPO}/contents/{repo_path}?ref={GITHUB_BRANCH}"
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "pastahouse-hr-rentals",
    }
    try:
        req = urllib.request.Request(url, headers=headers)
        with urllib.request.urlopen(req, timeout=20) as response:
            obj = json.loads(response.read().decode("utf-8"))
        return "found", base64.b64decode(obj["content"].replace("\n", ""))
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return "missing", None
        return "error", f"GitHub HTTP {e.code}"
    except Exception as e:
        return "error", str(e)


def _github_delete_file(repo_path, message):
    """Permanently delete a workbook in GitHub (also handles already deleted)."""
    token = _github_token()
    if not token:
        return False, "مفتاح github_token غير موجود في Streamlit Secrets"
    url = f"https://api.github.com/repos/{GITHUB_OWNER}/{GITHUB_REPO}/contents/{repo_path}"
    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "pastahouse-hr-rentals",
    }
    try:
        req = urllib.request.Request(f"{url}?ref={GITHUB_BRANCH}", headers=headers)
        with urllib.request.urlopen(req, timeout=20) as response:
            sha = json.loads(response.read().decode("utf-8"))["sha"]
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return True, ""
        return False, f"GitHub GET HTTP {e.code}"
    except Exception as e:
        return False, str(e)
    payload = json.dumps({"message": message, "sha": sha, "branch": GITHUB_BRANCH}).encode("utf-8")
    try:
        req = urllib.request.Request(url, data=payload, method="DELETE", headers={**headers, "Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=30) as response:
            if response.status not in (200, 204):
                return False, f"GitHub DELETE HTTP {response.status}"
        return True, ""
    except urllib.error.HTTPError as e:
        return False, f"GitHub DELETE HTTP {e.code}"
    except Exception as e:
        return False, str(e)


def save_df(key, df):
    """Persist in GitHub FIRST so a failed remote write isn't reported as saved."""
    import tempfile
    with tempfile.TemporaryDirectory() as directory:
        temporary_file = Path(directory) / FILES[key][1]
        df.to_excel(temporary_file, index=False)
        ok, err = _github_put_file(
            f"data/{FILES[key][1]}", temporary_file,
            f"Update {FILES[key][1]} from Streamlit"
        )
        if not ok:
            raise RuntimeError(f"تعذر الحفظ الدائم لملف {FILES[key][0]}: {err}")
        # The local copy is only a convenience; GitHub remains authoritative.
        p = DATA / FILES[key][1]
        p.write_bytes(temporary_file.read_bytes())
    _github_read_file.clear()
    return True


def delete_df(key):
    """Delete the remote file first; a redeploy cannot restore it."""
    ok, err = _github_delete_file(
        f"data/{FILES[key][1]}", f"Delete {FILES[key][1]} from Streamlit"
    )
    if not ok:
        raise RuntimeError(f"تعذر حذف ملف {FILES[key][0]} بشكل دائم: {err}")
    p = DATA / FILES[key][1]
    if p.exists():
        p.unlink()
    _github_read_file.clear()
    return True

def log_change(action, details):
    now = datetime.now()
    row = pd.DataFrame([{"التاريخ":now.strftime("%Y-%m-%d"),"الوقت":now.strftime("%H:%M:%S"),"العملية":action,"التفاصيل":details}])
    save_df("changes", pd.concat([load_df("changes"), row], ignore_index=True))

def find_col(df, names):
    mp = {str(c).strip().lower(): c for c in df.columns}
    for n in names:
        if n.strip().lower() in mp: return mp[n.strip().lower()]
    return None

def _clean_date_parts(value):
    """Return (year, month, day) from Excel serial/date/string values without forcing Hijri into Gregorian."""
    if pd.isna(value): return None
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return int(value.year), int(value.month), int(value.day)

    # Excel stores many Gregorian dates as serial numbers (for example 46419).
    # Convert them first so Bayan's rent sheet works without manual date cleanup.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            serial=float(value)
            if 20000 <= serial <= 80000:
                dt=pd.to_datetime(serial, unit="D", origin="1899-12-30")
                return int(dt.year), int(dt.month), int(dt.day)
        except Exception:
            pass

    text=str(value).strip()
    if not text or text.lower() in {"nan","none","nat"}: return None
    # Numeric Excel serials may also arrive as strings such as "46419" or "46419.0".
    try:
        serial=float(text)
        if 20000 <= serial <= 80000:
            dt=pd.to_datetime(serial, unit="D", origin="1899-12-30")
            return int(dt.year), int(dt.month), int(dt.day)
    except Exception:
        pass

    # Excel may return timestamps as strings; keep only the date part.
    text=text.split()[0].replace(".", "/").replace("-", "/")
    parts=text.split("/")
    if len(parts) < 3: return None
    try:
        a,b,c=[int(float(x)) for x in parts[:3]]
    except Exception:
        return None
    # yyyy/mm/dd or dd/mm/yyyy
    if a >= 1000: return a,b,c
    if c >= 1000: return c,b,a
    return None

def _is_hijri_year(year):
    return 1300 <= int(year) <= 1600

def _gregorian_jdn(y,m,d):
    a=(14-m)//12; yy=y+4800-a; mm=m+12*a-3
    return d+(153*mm+2)//5+365*yy+yy//4-yy//100+yy//400-32045

def _hijri_jdn(y,m,d):
    # Civil Islamic calendar calculation. We keep the user's original Hijri date in the UI.
    import math
    return (d + math.ceil(29.5*(m-1)) + (y-1)*354 + (3+11*y)//30 + 1948439 - 1)

def _value_date_info(value):
    parts=_clean_date_parts(value)
    if not parts: return None
    y,m,d=parts
    try:
        if _is_hijri_year(y):
            if not (1 <= m <= 12 and 1 <= d <= 30): return None
            target_jdn=_hijri_jdn(y,m,d)
            today=date.today()
            today_jdn=_gregorian_jdn(today.year,today.month,today.day)
            days=int(target_jdn-today_jdn)
            return {"kind":"هجري","days":days,"display":f"{y:04d}/{m:02d}/{d:02d}"}
        # Gregorian date: validate with datetime.date
        gd=date(y,m,d)
        return {"kind":"ميلادي","days":int((gd-date.today()).days),"display":f"{y:04d}/{m:02d}/{d:02d}"}
    except Exception:
        return None

def alert_table(key, date_names, days, status_names=None, unpaid_only=False):
    df=load_df(key)
    if key=="rents" and not df.empty:
        df=_prepare_rents_df(df)
        df,monthly_added=_ensure_monthly_rent_due(df)
        df=_refresh_rent_statuses(df)
        if monthly_added:
            save_df("rents",df)
        # The old blue-row date is only a monthly anchor/reference, not an unpaid bill.
        df=df[~df.apply(_is_monthly_reference,axis=1)].copy()
    if df.empty: return pd.DataFrame()
    dc=find_col(df,date_names)
    if not dc: return pd.DataFrame()
    x=df.copy()
    infos=x[dc].apply(_value_date_info)
    x["_days"]=infos.apply(lambda z: z["days"] if z else pd.NA)
    x["_date_kind"]=infos.apply(lambda z: z["kind"] if z else "")
    x["_display_date"]=infos.apply(lambda z: z["display"] if z else "")
    x["_date_col"]=dc
    x=x[x["_days"].notna()].copy()
    x["_days"]=x["_days"].astype(int)
    # Keep expired rows for separate expired alerts, plus only future rows inside the requested window.
    x=x[x["_days"]<=days]
    if unpaid_only:
        sc=find_col(x,status_names or [])
        if sc:
            x=x[~x[sc].astype(str).str.strip().str.lower().isin({"مسدد","تم السداد","paid","yes","نعم"})]
    return x

def _upcoming(df):
    if df.empty: return df
    return df[(df["_days"]>=0)].copy()

def _expired(df):
    if df.empty: return df
    return df[(df["_days"]<0)].copy()

def _first_value(row, names, fallback="غير محدد"):
    for n in names:
        if n in row.index and pd.notna(row[n]) and str(row[n]).strip() not in {"", "nan", "None"}:
            return str(row[n]).strip()
    return fallback

def _doc_alerts(df, doc_label, emoji, person_names=None):
    if df.empty:
        st.markdown('<div class="empty">لا توجد تنبيهات في هذا القسم</div>',unsafe_allow_html=True)
        return
    person_names=person_names or ["اسم الموظف","الموظف","الاسم","employee_name","name"]
    # Nearest/most recently expired first.
    for _,row in df.sort_values("_days", ascending=True).iterrows():
        who=_first_value(row,person_names)
        d=int(row["_days"]); dt=row["_display_date"]; kind=row["_date_kind"]
        if d < 0:
            st.error(f"🔴 {doc_label} منتهية: {doc_label} الموظف {who} انتهت منذ {abs(d)} يوم بتاريخ {dt} ({kind}).")
        elif d == 0:
            st.warning(f"{emoji} {doc_label}: {doc_label} الموظف {who} تنتهي اليوم بتاريخ {dt} ({kind}).")
        else:
            st.warning(f"{emoji} {doc_label}: {doc_label} الموظف {who} ستنتهي بتاريخ {dt} ({kind}) — متبقي {d} يومًا.")

def _format_installment_no(value):
    """Show installment numbers as 1, 2, 3... instead of Excel-style 1.0, 2.0."""
    if value is None or pd.isna(value): return ""
    text=str(value).strip()
    if not text or text.lower() in {"nan","none"}: return ""
    if text=="شهري": return "شهري"
    try:
        num=float(text)
        if num.is_integer(): return str(int(num))
    except Exception:
        pass
    return text

def _rent_alerts(df):
    if df.empty:
        st.markdown('<div class="empty">لا توجد تنبيهات إيجارات في هذا القسم</div>',unsafe_allow_html=True)
        return
    for _,row in df.sort_values("_days", ascending=True).iterrows():
        branch=_first_value(row,["الفرع","اسم الفرع","الموقع","اسم الموقع","branch","site"],"الفرع غير محدد")
        installment=_format_installment_no(_first_value(row,["رقم الدفعة","رقم الدفعه","installment_no"],""))
        if installment:
            installment_text=" — دفعة شهرية" if installment=="شهري" else f" — الدفعة {installment}"
        else:
            installment_text=""
        d=int(row["_days"]); dt=row["_display_date"]; kind=row["_date_kind"]
        if d < 0:
            st.error(f"🔴 إيجار متأخر: {branch}{installment_text} — متأخر {abs(d)} يوم — كان الاستحقاق {dt} ({kind}).")
        elif d == 0:
            st.warning(f"🟠 مستحق اليوم: {branch}{installment_text} — موعد السداد اليوم {dt} ({kind}).")
        else:
            st.warning(f"🟡 قرب موعد السداد: {branch}{installment_text} — متبقي {d} يوم — الاستحقاق {dt} ({kind}).")

def current_employees():
    df=load_df("employees")
    if df.empty:return df
    sc=find_col(df,["الحالة","status"])
    if sc:
        bad={"منتهي الخدمة","انتهت الخدمة","مستقيل","terminated","inactive","archive"}
        return df[~df[sc].astype(str).str.strip().str.lower().isin(bad)].copy()
    return df

def is_hr(): return st.session_state.get("role")=="hr"

def login():
    st.markdown('<div class="hero"><h2>نظام إدارة الموارد البشرية</h2><p>بيت المكرونة | Pasta House</p></div>',unsafe_allow_html=True)
    a,b,c=st.columns([1,1.2,1])
    with b:
        with st.form("login"):
            u=st.text_input("اسم المستخدم")
            p=st.text_input("كلمة المرور",type="password")
            ok=st.form_submit_button("تسجيل الدخول",use_container_width=True)
        if ok:
            rec=USERS.get(u.strip().lower())
            if rec and rec["password"]==hpwd(p):
                st.session_state.logged=True;st.session_state.role=rec["role"];st.session_state.display_name=rec["name"];st.rerun()
            else:st.error("بيانات الدخول غير صحيحة")

def global_search():
    q=st.text_input("🔍 البحث الشامل: موظف، رقم وظيفي، فرع، عقد، إقامة، كرت بلدي أو إيجار")
    if not q:return
    rows=[];ql=q.strip().lower()
    for key,(label,_) in FILES.items():
        if key in {"changes","movements"}:continue
        df=load_df(key)
        if df.empty:continue
        mask=df.astype(str).apply(lambda r:r.str.lower().str.contains(ql,na=False).any(),axis=1)
        hit=df[mask].copy()
        if not hit.empty:
            hit.insert(0,"القسم",label);rows.append(hit)
    if rows:st.dataframe(pd.concat(rows,ignore_index=True),use_container_width=True,hide_index=True)
    else:st.warning("لا توجد نتائج")

def dashboard():
    emp=current_employees()
    contracts_all=alert_table("contracts",["تاريخ نهاية العقد","تاريخ انتهاء العقد","contract_end"],90)
    iqama_all=alert_table("iqama",["تاريخ انتهاء الإقامة","انتهاء الإقامة","تاريخ انتهاء الاقامة","iqama_end"],90)
    passports_all=alert_table("iqama",["تاريخ انتهاء الجواز","انتهاء الجواز","تاريخ نهاية الجواز","passport_end"],90)
    baladi_all=alert_table("baladi",["تاريخ انتهاء الكرت البلدي","تاريخ انتهاء الكرت","انتهاء الكرت البلدي","تاريخ انتهاء البطاقة البلدية","تاريخ الانتهاء","baladi_end"],20)
    rents_all=alert_table("rents",["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"],ALERT_DAYS["rents"],["حالة السداد","status"],True)

    contracts=_upcoming(contracts_all); iqama=_upcoming(iqama_all); passports=_upcoming(passports_all)
    baladi=_upcoming(baladi_all); due=_upcoming(rents_all); overdue=_expired(rents_all)

    st.markdown('<div class="hero"><h2>لوحة التحكم</h2><p>نظرة عامة على الموظفين والوثائق والفروع والإيجارات</p></div>',unsafe_allow_html=True)
    global_search()
    c=st.columns(2);c[0].metric("إجمالي الموظفين",len(emp));c[1].metric("أرشيف الموظفين",len(load_df("archive")))
    c=st.columns(5);c[0].metric("عقود خلال 3 أشهر",len(contracts));c[1].metric("إقامات خلال 3 أشهر",len(iqama));c[2].metric("جوازات خلال 3 أشهر",len(passports));c[3].metric("كروت بلدية خلال 20 يوم",len(baladi));c[4].metric("إيجارات مستحقة قريبًا",len(due))
    st.metric("🚨 إيجارات متأخرة",len(overdue))

    st.subheader("🚨 مركز التنبيهات")
    tabs=st.tabs(["الإيجارات المتأخرة","الإيجارات المستحقة","العقود","الإقامات","الجوازات","الكروت البلدية"])
    with tabs[0]: _rent_alerts(overdue)
    with tabs[1]: _rent_alerts(due)
    with tabs[2]: _doc_alerts(contracts_all,"العقد","📄")
    with tabs[3]: _doc_alerts(iqama_all,"الإقامة","🪪")
    with tabs[4]: _doc_alerts(passports_all,"الجواز","🛂")
    with tabs[5]: _doc_alerts(baladi_all,"الكرت البلدي","🏥")

    st.subheader("آخر التعديلات")
    ch=load_df("changes")
    if ch.empty:st.markdown('<div class="empty">لا توجد تعديلات مسجلة</div>',unsafe_allow_html=True)
    else:st.dataframe(ch.tail(10).iloc[::-1],use_container_width=True,hide_index=True)

def data_page(key,title):
    st.title(title);df=load_df(key)
    if df.empty:st.markdown('<div class="empty">لا توجد بيانات بعد. ارفعي الملف من إدارة البيانات.</div>',unsafe_allow_html=True);return
    q=st.text_input("بحث داخل الجدول");view=df
    if q:view=df[df.astype(str).apply(lambda r:r.str.contains(q,case=False,na=False).any(),axis=1)]
    if is_hr():
        edited=st.data_editor(view,use_container_width=True,num_rows="dynamic")
        if st.button("💾 حفظ التعديلات",type="primary"):
            out=edited if len(view)==len(df) and list(view.index)==list(df.index) else df.copy()
            if out is not edited:out.loc[edited.index,edited.columns]=edited
            save_df(key,out);log_change("تعديل بيانات",title);st.success("تم الحفظ");st.rerun()
    else:st.dataframe(view,use_container_width=True,hide_index=True)

def add_employee():
    st.title("➕ إضافة موظف")
    if not is_hr():st.warning("حساب المدير للمشاهدة فقط");return
    with st.form("add_emp"):
        a,b=st.columns(2)
        emp_id=a.text_input("الرقم الوظيفي *");name=b.text_input("اسم الموظف *")
        branch=a.text_input("الفرع / الموقع *");job=b.text_input("المسمى الوظيفي")
        nationality=a.text_input("الجنسية");start=b.date_input("تاريخ المباشرة")
        contract_end=a.date_input("تاريخ نهاية العقد");iqama_end=b.date_input("تاريخ انتهاء الإقامة")
        baladi_end=a.date_input("تاريخ انتهاء الكرت البلدي")
        ok=st.form_submit_button("حفظ الموظف",use_container_width=True)
    if ok:
        if not emp_id or not name or not branch:st.error("الرقم الوظيفي والاسم والفرع مطلوبة");return
        emp=load_df("employees")
        save_df("employees",pd.concat([emp,pd.DataFrame([{"الرقم الوظيفي":emp_id,"اسم الموظف":name,"الفرع الحالي":branch,"المسمى الوظيفي":job,"الجنسية":nationality,"تاريخ المباشرة":start,"الحالة":"على رأس العمل"}])],ignore_index=True))
        for key,data in [("contracts",{"الرقم الوظيفي":emp_id,"اسم الموظف":name,"تاريخ نهاية العقد":contract_end}),("iqama",{"الرقم الوظيفي":emp_id,"اسم الموظف":name,"تاريخ انتهاء الإقامة":iqama_end}),("baladi",{"الرقم الوظيفي":emp_id,"اسم الموظف":name,"تاريخ انتهاء الكرت البلدي":baladi_end})]:
            save_df(key,pd.concat([load_df(key),pd.DataFrame([data])],ignore_index=True))
        log_change("إضافة موظف",f"{name} - {emp_id}");st.success("تمت إضافة الموظف");st.rerun()

def move_employee():
    st.title("📜 نقل موظف / سجل الحركات")
    if not is_hr():data_page("movements","سجل الحركات");return
    emp=current_employees();idc=find_col(emp,["الرقم الوظيفي","employee_id"]);nc=find_col(emp,["اسم الموظف","employee_name"]);bc=find_col(emp,["الفرع الحالي","الفرع","branch"])
    if emp.empty or not idc:st.info("ارفعي ملف الموظفين أولًا");return
    labels=[f"{r[idc]} - {r[nc] if nc else ''}" for _,r in emp.iterrows()]
    selected=st.selectbox("الموظف",labels);row=emp.iloc[labels.index(selected)];old_branch=str(row[bc]) if bc else ""
    new_branch=st.text_input("الفرع / الموقع الجديد");move_date=st.date_input("تاريخ النقل");note=st.text_input("ملاحظة")
    if st.button("تنفيذ النقل",type="primary"):
        if not new_branch:st.error("اكتبي الموقع الجديد");return
        full=load_df("employees");fid=find_col(full,["الرقم الوظيفي","employee_id"]);fbc=find_col(full,["الفرع الحالي","الفرع","branch"])
        full.loc[full[fid].astype(str)==str(row[idc]),fbc]=new_branch;save_df("employees",full)
        rec={"التاريخ":move_date,"الرقم الوظيفي":row[idc],"اسم الموظف":row[nc] if nc else "","من":old_branch,"إلى":new_branch,"ملاحظة":note}
        save_df("movements",pd.concat([load_df("movements"),pd.DataFrame([rec])],ignore_index=True))
        log_change("نقل موظف",f"{row[nc] if nc else row[idc]} من {old_branch} إلى {new_branch}");st.success("تم النقل");st.rerun()
    st.markdown("---");data_page("movements","سجل الحركات")

def archive_employee():
    st.title("📂 إنهاء خدمة / أرشيف الموظفين")
    if not is_hr():data_page("archive","أرشيف الموظفين");return
    emp=current_employees();idc=find_col(emp,["الرقم الوظيفي","employee_id"]);nc=find_col(emp,["اسم الموظف","employee_name"])
    if emp.empty or not idc:st.info("لا توجد بيانات موظفين");return
    labels=[f"{r[idc]} - {r[nc] if nc else ''}" for _,r in emp.iterrows()]
    row=emp.iloc[labels.index(st.selectbox("اختاري الموظف",labels))]
    reason=st.selectbox("سبب إنهاء الخدمة",["انتهاء عقد","استقالة","إنهاء خدمة","أخرى"]);end_date=st.date_input("تاريخ نهاية الخدمة")
    if st.button("نقل إلى الأرشيف",type="primary"):
        full=load_df("employees");fid=find_col(full,["الرقم الوظيفي","employee_id"]);match=full[fid].astype(str)==str(row[idc])
        archived=full[match].copy();archived["تاريخ نهاية الخدمة"]=end_date;archived["سبب نهاية الخدمة"]=reason
        save_df("archive",pd.concat([load_df("archive"),archived],ignore_index=True));save_df("employees",full[~match].copy())
        log_change("إنهاء خدمة",f"{row[nc] if nc else row[idc]} - {reason}");st.success("تم نقل الموظف إلى الأرشيف");st.rerun()
    st.markdown("---");data_page("archive","أرشيف الموظفين")


IGNORED_RENT_SITES = {"سكن نادر", "مكتب البلدية"}

def _drop_ignored_rent_rows(df):
    """Ignore the source workbook rows Bayan intentionally keeps hidden."""
    if df is None or df.empty:
        return df
    x=df.copy()
    site_col=find_col(x,["الموقع","اسم الموقع","الفرع","اسم الفرع","branch","site"])
    if not site_col:
        return x
    site_text=x[site_col].astype(str).str.strip()
    # These two rows are hidden in the supplied rent workbook and should never
    # participate in display, alerts, delay calculations, or payment actions.
    x=x[~site_text.isin(IGNORED_RENT_SITES)].copy()
    # Also suppress placeholder rows that contain only a site name and no real data.
    data_cols=[c for c in x.columns if c != site_col]
    if data_cols:
        def has_real_data(row):
            for c in data_cols:
                v=row.get(c)
                if pd.isna(v):
                    continue
                t=str(v).strip()
                if t and t.lower() not in {"none","nan","nat"}:
                    return True
            return False
        x=x[x.apply(has_real_data,axis=1)].copy()
    return x

def _normalize_rent_status_text(value):
    text=str(value).strip() if pd.notna(value) else ""
    low=text.lower().replace("ة","ه")
    paid={"مسدد","تم السداد","paid","yes","نعم"}
    if text.lower() in paid: return "مسدد"
    if low in {"غير مستحقه","غير مسحقه","غير مستحق","غير مسحق"}: return "غير مستحق"
    if text in {"يستحق قريب","مستحق قريب"}: return "يستحق قريب"
    if text in {"غير مسدد","غير مدفوع"}: return "غير مسدد"
    return text


def _rent_col_key(value):
    """Arabic-insensitive key used to recognize Bayan's different rent-sheet layouts."""
    text=str(value).strip().lower()
    for a,b in [("أ","ا"),("إ","ا"),("آ","ا"),("ة","ه"),("ى","ي")]:
        text=text.replace(a,b)
    return " ".join(text.split())


def _rent_due_columns(df):
    """Return wide-layout due-date columns in installment order (1..4)."""
    found={}
    for c in df.columns:
        k=_rent_col_key(c)
        # Screenshot layout: الاستحقاق الاول / الثاني / الثالث / الرابع
        if "الاستحقاق" not in k and "استحقاق" not in k:
            continue
        if any(w in k for w in ["الاول","الاولى","1"]): found[1]=c
        elif any(w in k for w in ["الثاني","الثانيه","2"]): found[2]=c
        elif any(w in k for w in ["الثالث","الثالثه","3"]): found[3]=c
        elif any(w in k for w in ["الرابع","الرابعه","4"]): found[4]=c
    return [found[n] for n in sorted(found)]


def _expand_wide_rent_sheet(df):
    """
    Convert the visual Excel layout (one site per row + 2/4 due-date columns)
    into the dashboard's internal one-installment-per-row layout.
    Blank third/fourth cells mean the contract has only two installments.
    The monthly blue row stays monthly and is handled by _ensure_monthly_rent_due.
    """
    # Bayan also has an older/normalized file where every payment is already a row
    # and legacy columns such as "تاريخ الاستحقاق الثاني" still exist. If there is
    # a payment-number column + a main due-date column, it is already normalized.
    number_col=find_col(df,["رقم الدفعة","رقم الدفعه","installment_no"])
    main_due=find_col(df,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
    if number_col and main_due:
        return df

    due_cols=_rent_due_columns(df)
    if len(due_cols) < 2:
        return df

    site_col=find_col(df,["الموقع","اسم الموقع","الفرع","اسم الفرع","branch","site"])
    rent_col=find_col(df,["الإيجار","الايجار","قيمة الإيجار","قيمة الايجار","rent"])
    pay_col=find_col(df,["الدفعة","مبلغ الدفعة","payment"])
    status_col=find_col(df,["حالة السداد","الحالة","status"])

    # Columns that may contain notes merged across the due-date area in the source sheet.
    rows=[]
    for _,r in df.iterrows():
        site=str(r.get(site_col,"")).strip() if site_col else ""
        if not site or site.lower() in {"nan","none"}:
            continue

        combined=" ".join(str(r.get(c,"")) for c in df.columns if pd.notna(r.get(c,"")))
        monthly=("شهري" in combined.replace("شهرياً","شهري").replace("شهريا","شهري") and
                 ("كل دفعه" in combined or "كل دفعة" in combined))

        if monthly:
            # First real date is the monthly anchor. Note text may occupy later due cells.
            anchor=None
            for c in due_cols:
                if _value_date_info(r.get(c)):
                    anchor=r.get(c); break
            if anchor is None:
                continue
            amount=1100
            # Prefer the explicit monthly amount from the source text if present.
            import re
            m=re.search(r"(?:دفعه|دفعة)\s*([0-9,]+)\s*ريال",combined)
            if m:
                try: amount=float(m.group(1).replace(",",""))
                except Exception: pass
            elif rent_col and pd.notna(r.get(rent_col)):
                try: amount=float(r.get(rent_col))
                except Exception: pass
            rows.append({
                "الموقع":site,
                "الإيجار":r.get(rent_col,"") if rent_col else "",
                "الدفعة":amount,
                "رقم الدفعة":"شهري",
                "تاريخ الاستحقاق":anchor,
                "حالة السداد":"مرجع شهري",
                "نوع السجل":"مرجع شهري",
                "ملاحظة":"الإيجار شهري كل دفعة 1100 ريال" if "1100" in combined else "الإيجار شهري",
            })
            continue

        # Normal contracts: every nonblank/valid date cell is one installment.
        installment_no=0
        for c in due_cols:
            value=r.get(c)
            if not _value_date_info(value):
                # Empty third/fourth columns = no extra installment, not an error.
                continue
            installment_no+=1
            rows.append({
                "الموقع":site,
                "الإيجار":r.get(rent_col,"") if rent_col else "",
                "الدفعة":r.get(pay_col,"") if pay_col else "",
                "رقم الدفعة":installment_no,
                "تاريخ الاستحقاق":value,
                "حالة السداد":"",
                "نوع السجل":"دفعة",
                "ملاحظة":"",
            })

    return pd.DataFrame(rows) if rows else df


def _prepare_rents_df(df):
    """Normalize Bayan rent sheet; accepts both row-per-installment and 2/4-column layouts."""
    if df is None or df.empty: return df
    x=df.copy()
    x.columns=[str(c).strip() for c in x.columns]
    x=_drop_ignored_rent_rows(x)
    if x.empty:
        return x

    # If Bayan uploads the screenshot-style sheet, expand it internally first.
    x=_expand_wide_rent_sheet(x)
    x.columns=[str(c).strip() for c in x.columns]

    rename={}
    for c in x.columns:
        k=_rent_col_key(c)
        if k=="الايجار": rename[c]="الإيجار"
        elif k in {"رقم الدفعه","رقم الدفعه"}: rename[c]="رقم الدفعة"
        elif k in {"تاريخ الاستحقاق","تاريخ استحقاق السداد"}: rename[c]="تاريخ الاستحقاق"
        elif k=="حاله السداد": rename[c]="حالة السداد"
    x=x.rename(columns=rename)

    if "حالة السداد" not in x.columns: x["حالة السداد"]=""
    x["حالة السداد"]=x["حالة السداد"].apply(_normalize_rent_status_text)
    if "نوع السجل" not in x.columns: x["نوع السجل"]=""
    if "تاريخ السداد الفعلي" not in x.columns: x["تاريخ السداد الفعلي"]=""
    if "تفاصيل السداد" not in x.columns: x["تفاصيل السداد"]=""
    if "تاريخ الاستحقاق الأصلي" not in x.columns:
        dc=find_col(x,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
        x["تاريخ الاستحقاق الأصلي"]=x[dc] if dc else ""
    if "تاريخ آخر تعديل للاستحقاق" not in x.columns: x["تاريخ آخر تعديل للاستحقاق"]=""
    if "ملاحظة تعديل الاستحقاق" not in x.columns: x["ملاحظة تعديل الاستحقاق"]=""
    x=_mark_legacy_monthly_reference(x)
    return x


def _monthly_rent_row(row):
    text=" ".join(str(v) for v in row.values if pd.notna(v)).replace("شهرياً","شهري").replace("شهريا","شهري")
    return "شهري" in text and ("كل دفعه" in text or "كل دفعة" in text)


def _add_months(d, months=1):
    import calendar
    m=d.month-1+months; y=d.year+m//12; m=m%12+1
    return date(y,m,min(d.day,calendar.monthrange(y,m)[1]))


def _date_from_value(value):
    parts=_clean_date_parts(value)
    if not parts: return None
    y,m,d=parts
    if _is_hijri_year(y): return None
    try:return date(y,m,d)
    except Exception:return None


def _rent_is_paid(value):
    return _normalize_rent_status_text(value)=="مسدد"


def _is_monthly_reference(row):
    record_type=str(row.get("نوع السجل","")).strip()
    status=str(row.get("حالة السداد","")).strip()
    return record_type=="مرجع شهري" or status=="مرجع شهري"


def _mark_legacy_monthly_reference(df):
    """Recognize the original old blue monthly row in previously-saved rent data."""
    if df is None or df.empty:return df
    x=df.copy()
    if "نوع السجل" not in x.columns:x["نوع السجل"]=""
    branch_col=find_col(x,["اسم الفرع","الفرع","الموقع","اسم الموقع","branch"])
    due_col=find_col(x,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
    paid_col="تاريخ السداد الفعلي" if "تاريخ السداد الفعلي" in x.columns else None
    if not branch_col or not due_col:return x
    for branch,grp in x.groupby(x[branch_col].astype(str).str.strip()):
        mids=[idx for idx,row in grp.iterrows() if _monthly_rent_row(row)]
        if not mids:continue
        dated=[]
        for idx in mids:
            gd=_date_from_value(x.at[idx,due_col])
            if gd:dated.append((idx,gd))
        if not dated:continue
        idx,gd=min(dated,key=lambda z:z[1])
        paid=bool(paid_col and str(x.at[idx,paid_col]).strip() not in {"","nan","None"})
        # The source sheet's blue anchor is years old (2022). Do not treat it as arrears.
        if not paid and (date.today()-gd).days>180:
            x.at[idx,"نوع السجل"]="مرجع شهري"
            x.at[idx,"حالة السداد"]="مرجع شهري"
    return x


def _rent_live_status(due_value, paid=False):
    if paid: return "مسدد"
    info=_value_date_info(due_value)
    if not info: return "غير محدد"
    d=int(info["days"])
    if d < 0: return f"متأخر {abs(d)} يوم"
    if d == 0: return "مستحق اليوم"
    if d <= ALERT_DAYS["rents"]: return f"قرب موعد السداد — متبقي {d} يوم"
    return "غير مستحق"


def _refresh_rent_statuses(df):
    if df is None or df.empty:return df
    x=df.copy()
    sc=find_col(x,["حالة السداد","الحالة","status"])
    dc=find_col(x,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
    if not sc or not dc:return x
    for idx,row in x.iterrows():
        if _is_monthly_reference(row):
            x.at[idx,sc]="مرجع شهري"
        elif _rent_is_paid(row.get(sc)):
            x.at[idx,sc]="مسدد"
        else:
            x.at[idx,sc]=_rent_live_status(row.get(dc),False)
    return x


def _ensure_monthly_rent_due(df):
    """Keep paid monthly rows as history and create one current/next unpaid monthly installment when needed."""
    if df is None or df.empty:return df,False
    branch_col=find_col(df,["اسم الفرع","الفرع","الموقع","اسم الموقع","branch"])
    due_col=find_col(df,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
    status_col=find_col(df,["حالة السداد","الحالة","status"])
    payment_col=find_col(df,["الدفعة","مبلغ الدفعة","payment"])
    if not branch_col or not due_col or not status_col:return df,False

    x=df.copy(); changed=False; today=date.today()
    monthly_rows=[idx for idx,row in x.iterrows() if _monthly_rent_row(row)]
    branches=[]
    for idx in monthly_rows:
        b=str(x.at[idx,branch_col]).strip()
        if b and b not in branches:branches.append(b)

    for branch in branches:
        group=x[x[branch_col].astype(str).str.strip()==branch]
        anchors=[]
        for idx,row in group.iterrows():
            if _monthly_rent_row(row):
                gd=_date_from_value(row[due_col])
                if gd:anchors.append((idx,gd))
        if not anchors:continue
        anchor_idx,anchor=max(anchors,key=lambda z:z[1])
        try:target=date(today.year,today.month,min(anchor.day,28 if today.month==2 else anchor.day))
        except Exception:target=date(today.year,today.month,1)

        # If this month's installment already exists and is paid, prepare next month instead.
        existing_dates=[]
        for idx,row in group.iterrows():
            gd=_date_from_value(row[due_col])
            if gd:existing_dates.append((idx,gd,_rent_is_paid(row[status_col])))
        same=[z for z in existing_dates if z[1]==target]
        if same and any(z[2] for z in same): target=_add_months(target,1)
        if any(z[1]==target and not z[2] for z in existing_dates): continue

        newrow=x.loc[anchor_idx].copy()
        newrow[due_col]=target
        newrow["تاريخ الاستحقاق الأصلي"]=target
        newrow["نوع السجل"]="دفعة شهرية"
        newrow[status_col]=_rent_live_status(target,False)
        newrow["تاريخ السداد الفعلي"]=""
        newrow["تفاصيل السداد"]=""
        newrow["تاريخ آخر تعديل للاستحقاق"]=""
        newrow["ملاحظة تعديل الاستحقاق"]=""
        if payment_col:newrow[payment_col]=1100
        x=pd.concat([x,pd.DataFrame([newrow])],ignore_index=True)
        changed=True
    return x,changed


def _rent_display_df(df):
    """User-facing rent table: hide internal helper fields, blanks instead of None, clean dates."""
    if df is None or df.empty:
        return df
    x=_drop_ignored_rent_rows(df.copy())
    # Do not show the monthly anchor/reference row; it is only used internally.
    if "نوع السجل" in x.columns:
        x=x[x["نوع السجل"].astype(str).str.strip()!="مرجع شهري"].copy()

    wanted=[
        "الموقع","اسم الموقع","الفرع","اسم الفرع",
        "الإيجار","الدفعة","رقم الدفعة","تاريخ الاستحقاق",
        "حالة السداد","تاريخ السداد الفعلي","تفاصيل السداد"
    ]
    cols=[]
    for c in wanted:
        if c in x.columns and c not in cols:
            cols.append(c)
    if cols:
        x=x[cols].copy()

    def clean(v):
        if pd.isna(v): return ""
        if isinstance(v,(pd.Timestamp,datetime,date)):
            return f"{v.year:04d}/{v.month:02d}/{v.day:02d}"
        text=str(v).strip()
        if text.lower() in {"none","nan","nat"}: return ""
        # Remove Excel/Pandas midnight timestamp from otherwise normal dates.
        if text.endswith(" 00:00:00"):
            text=text[:-9]
        return text.replace("-","/") if len(text)>=10 and text[:4].isdigit() and text[4] in "-/" else text

    for c in x.columns:
        x[c]=x[c].map(clean)
    return x


def rents_page():
    st.title("🏠 الإيجارات")
    df=_prepare_rents_df(load_df("rents")).reset_index(drop=True)
    if df.empty:
        st.markdown('<div class="empty">ارفعي ملف الإيجارات من إدارة البيانات</div>',unsafe_allow_html=True)
        return

    branch_col=find_col(df,["اسم الفرع","الفرع","الموقع","اسم الموقع","branch"])
    status_col=find_col(df,["حالة السداد","الحالة","status"])
    due_col=find_col(df,["تاريخ الاستحقاق","تاريخ استحقاق السداد","rent_due"])
    payment_col=find_col(df,["الدفعة","مبلغ الدفعة","payment"])
    number_col=find_col(df,["رقم الدفعة","رقم الدفعه","installment_no"])
    paid_date_col="تاريخ السداد الفعلي"

    # Preserve old monthly payments and create a separate current/next monthly installment.
    df,monthly_added=_ensure_monthly_rent_due(df)
    df=df.reset_index(drop=True)
    df=_refresh_rent_statuses(df).reset_index(drop=True)

    # Pandas may infer Excel date columns as datetime64. Streamlit edits and
    # payment actions write text values (YYYY/MM/DD or blank), so keep every
    # editable rent field as object to avoid LossySetitemError/TypeError.
    editable_rent_cols=[
        status_col, due_col, paid_date_col, "تفاصيل السداد",
        "تاريخ الاستحقاق الأصلي", "تاريخ آخر تعديل للاستحقاق",
        "ملاحظة تعديل الاستحقاق"
    ]
    for c in editable_rent_cols:
        if c and c in df.columns:
            df[c]=df[c].astype("object")

    if monthly_added:
        save_df("rents",df)

    if not is_hr():
        st.dataframe(_rent_display_df(df),use_container_width=True,hide_index=True)
        return

    st.caption("النظام يتعرف تلقائيًا على عقد بدفعتين أو 4 دفعات؛ الخانات الفارغة تُتجاهل. غير مستحق = موعده لم يأتِ بعد، والإيجار الشهري يُتابَع شهرًا بشهر.")

    if branch_col and status_col and due_col:
        st.subheader("✅ تسجيل سداد إيجار")
        options=[]
        for index,row in df.iterrows():
            if _is_monthly_reference(row) or _rent_is_paid(row.get(status_col)):continue
            branch=str(row.get(branch_col,""))
            due=_value_date_info(row.get(due_col))
            due_txt=due["display"] if due else str(row.get(due_col,""))
            n=""
            if number_col and pd.notna(row.get(number_col)):
                inst=_format_installment_no(row[number_col])
                if inst: n=" — دفعة شهرية" if inst=="شهري" else f" — الدفعة {inst}"
            amount=""
            if payment_col and pd.notna(row.get(payment_col)):
                try:amount=f" — {float(row[payment_col]):,.0f} ريال"
                except:pass
            options.append((index,f"{branch}{n}{amount} — استحقاق {due_txt}"))

        if options:
            labels=[x[1] for x in options]
            selected_label=st.selectbox("اختاري الموقع / الدفعة",labels,key="rent_pay_select")
            selected_index=dict((label,idx) for idx,label in options)[selected_label]
            row=df.loc[selected_index]
            info=_value_date_info(row[due_col])
            if info:
                d=int(info["days"])
                if d<0: st.error(f"🔴 متأخر عن الاستحقاق بـ {abs(d)} يوم")
                elif d==0: st.warning("🟠 الاستحقاق اليوم")
                elif d<=ALERT_DAYS["rents"]: st.warning(f"🟡 متبقي على الاستحقاق {d} يوم")
                else: st.info(f"🔵 غير مستحق — متبقي {d} يوم")

            if st.button("✅ تم السداد",type="primary",use_container_width=True):
                today=date.today(); original_due=row[due_col]
                delay_text=""
                if info:
                    delta=-int(info["days"])
                    if delta>0: delay_text=f"تم السداد — متأخر {delta} يوم"
                    elif delta==0: delay_text="تم السداد في الموعد"
                    else: delay_text=f"تم السداد قبل الاستحقاق بـ {abs(delta)} يوم"
                else: delay_text="تم السداد"
                # Ensure payment fields exist, then update exactly one row.
                if paid_date_col not in df.columns:
                    df[paid_date_col]=""
                if "تفاصيل السداد" not in df.columns:
                    df["تفاصيل السداد"]=""
                selected_index=int(selected_index)
                df.at[selected_index,status_col]="مسدد"
                df.at[selected_index,paid_date_col]=today.strftime("%Y/%m/%d")
                df.at[selected_index,"تفاصيل السداد"]=delay_text
                branch=str(df.at[selected_index,branch_col]).strip()
                log_change("سداد إيجار",f"{branch} — الاستحقاق {info['display'] if info else original_due} — السداد {today:%Y/%m/%d} — {delay_text}")
                save_df("rents",df)
                st.success(delay_text)
                st.rerun()
        else:
            st.success("لا توجد دفعات غير مسددة حاليًا")

        # Allow HR to undo a payment that was marked by mistake.
        paid_options=[]
        for index,row in df.iterrows():
            if _is_monthly_reference(row) or not _rent_is_paid(row.get(status_col)):
                continue
            branch=str(row.get(branch_col,""))
            due=_value_date_info(row.get(due_col))
            due_txt=due["display"] if due else str(row.get(due_col,""))
            n=""
            if number_col and pd.notna(row.get(number_col)):
                inst=_format_installment_no(row[number_col])
                if inst: n=" — دفعة شهرية" if inst=="شهري" else f" — الدفعة {inst}"
            amount=""
            if payment_col and pd.notna(row.get(payment_col)):
                try: amount=f" — {float(row[payment_col]):,.0f} ريال"
                except: pass
            paid_options.append((index,f"{branch}{n}{amount} — استحقاق {due_txt}"))

        if paid_options:
            with st.expander("↩️ إلغاء سداد مسجل بالخطأ"):
                paid_labels=[x[1] for x in paid_options]
                paid_label=st.selectbox("اختاري الدفعة المسددة",paid_labels,key="rent_unpay_select")
                paid_index=dict((label,idx) for idx,label in paid_options)[paid_label]
                if st.button("↩️ إلغاء السداد وإرجاعها غير مسددة",use_container_width=True):
                    undo_row=df.loc[paid_index]
                    undo_due=undo_row.get(due_col)
                    undo_info=_value_date_info(undo_due)
                    branch=str(undo_row.get(branch_col,"")).strip()
                    paid_index=int(paid_index)
                    df.at[paid_index,status_col]=_rent_live_status(undo_due,False)
                    if paid_date_col in df.columns:
                        df.at[paid_index,paid_date_col]=""
                    if "تفاصيل السداد" in df.columns:
                        df.at[paid_index,"تفاصيل السداد"]=""
                    log_change("إلغاء سداد إيجار",f"{branch} — الاستحقاق {undo_info['display'] if undo_info else undo_due} — أعيدت إلى غير مسدد")
                    save_df("rents",df)
                    st.success("تم إلغاء السداد وإرجاع الدفعة إلى غير مسددة")
                    st.rerun()

        st.markdown("---")
        st.subheader("📅 تعديل موعد استحقاق دفعة محددة")
        st.caption("هذه الصلاحية لحساب HR فقط. التعديل يغيّر الدفعة المختارة فقط ويحفظ تاريخها الأصلي.")
        edit_options=[]
        for index,row in df.iterrows():
            if _is_monthly_reference(row):continue
            branch=str(row.get(branch_col,""))
            due=_value_date_info(row.get(due_col)); due_txt=due["display"] if due else str(row.get(due_col,""))
            n=""
            if number_col and pd.notna(row.get(number_col)):
                inst=_format_installment_no(row[number_col])
                if inst: n=" — دفعة شهرية" if inst=="شهري" else f" — الدفعة {inst}"
            edit_options.append((index,f"{branch}{n} — {due_txt}"))
        edit_labels=[x[1] for x in edit_options]
        edit_label=st.selectbox("اختاري الدفعة التي تريدين تعديل موعدها",edit_labels,key="rent_edit_select")
        edit_idx=dict((label,idx) for idx,label in edit_options)[edit_label]
        old_due=df.at[edit_idx,due_col]
        old_info=_value_date_info(old_due)
        default_due=old_info["display"] if old_info else str(old_due)
        with st.form("rent_due_edit_form"):
            new_due=st.text_input("موعد الاستحقاق الجديد",value=default_due,help="مثال: 2026/10/20 أو 1448/04/08")
            note=st.text_input("سبب / ملاحظة التعديل (اختياري)")
            save_due=st.form_submit_button("💾 حفظ موعد الاستحقاق الجديد",use_container_width=True)
        if save_due:
            new_info=_value_date_info(new_due.strip())
            if not new_info:
                st.error("التاريخ غير واضح. اكتبيه مثل 2026/10/20 أو 1448/04/08")
            else:
                original_col="تاريخ الاستحقاق الأصلي"
                current_original=df.at[edit_idx,original_col] if original_col in df.columns else ""
                if pd.isna(current_original) or str(current_original).strip() in {"","nan","None"}:
                    df.at[edit_idx,original_col]=old_due
                df.at[edit_idx,due_col]=new_due.strip()
                df.at[edit_idx,"تاريخ آخر تعديل للاستحقاق"]=date.today().strftime("%Y/%m/%d")
                df.at[edit_idx,"ملاحظة تعديل الاستحقاق"]=note.strip()
                if not _rent_is_paid(df.at[edit_idx,status_col]):
                    df.at[edit_idx,status_col]=_rent_live_status(new_due.strip(),False)
                branch=str(df.at[edit_idx,branch_col]).strip()
                log_change("تعديل موعد إيجار",f"{branch} — من {default_due} إلى {new_info['display']}" + (f" — {note.strip()}" if note.strip() else ""))
                save_df("rents",df)
                st.success(f"تم تعديل موعد هذه الدفعة فقط إلى {new_info['display']}")
                st.rerun()
    else:
        st.warning("ملف الإيجارات يحتاج أعمدة: الموقع، تاريخ الاستحقاق، وحالة السداد.")

    st.markdown("---")
    st.subheader("جدول الإيجارات")
    st.caption("هذا هو العرض المرتب للمستخدم. الأعمدة الداخلية الخاصة بالحساب مخفية، والخانات الفارغة تظهر فارغة بدل None.")
    st.dataframe(_rent_display_df(df),use_container_width=True,hide_index=True)
    st.info("لتعديل موعد دفعة استخدمي قسم «تعديل موعد استحقاق دفعة محددة» بالأعلى؛ بهذه الطريقة لا يتأثر باقي العقد.")

def _read_uploaded_excel(f, key):
    """Read uploads; for rents, ignore rows hidden in the original XLSX workbook."""
    if key != "rents":
        return pd.read_excel(f)
    # Preserve upload bytes because openpyxl and pandas both need to read them.
    from io import BytesIO
    raw=f.getvalue() if hasattr(f,"getvalue") else f.read()
    hidden_excel_rows=set()
    try:
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(raw),data_only=True,read_only=False)
        ws=wb.active
        hidden_excel_rows={i for i in range(2,ws.max_row+1) if ws.row_dimensions[i].hidden}
    except Exception:
        hidden_excel_rows=set()
    incoming=pd.read_excel(BytesIO(raw))
    if hidden_excel_rows:
        # pandas row 0 corresponds to Excel row 2 (row 1 is the header).
        drop_idx=[r-2 for r in hidden_excel_rows if 0 <= r-2 < len(incoming)]
        incoming=incoming.drop(index=drop_idx).reset_index(drop=True)
    return incoming

def uploads():
    st.title("📤 إدارة البيانات")
    if not is_hr():st.warning("حساب المدير للمشاهدة فقط");return
    st.info("ارفعي كل ملف في قسمه. يُحفظ التحديث والحذف في GitHub بشكل دائم، ولا تعود النسخ القديمة بعد إعادة تشغيل المنصة.")
    keys=["employees","contracts","iqama","baladi","rents"];cols=st.columns(2)
    for i,key in enumerate(keys):
        with cols[i%2]:
            label=FILES[key][0];st.markdown(f"### {label}")
            f=st.file_uploader(f"رفع ملف {label}",type=["xlsx","xls"],key=f"up_{key}",help="XLSX, XLS — 200MB per file")
            cur=load_df(key);st.caption("لم يُرفع بعد" if cur.empty else f"آخر نسخة: {len(cur)} سجل")
            b1,b2=st.columns(2)
            with b1:
                if f and st.button(f"🔄 تحديث {label}",key=f"btn_{key}",use_container_width=True):
                    try:
                        incoming=_read_uploaded_excel(f,key)
                        if key=="rents": incoming=_refresh_rent_statuses(_prepare_rents_df(incoming))
                        save_df(key,incoming);log_change("رفع ملف",label);st.success(f"تم تحديث {label}");st.rerun()
                    except Exception as e:st.error(f"تعذر قراءة الملف: {e}")
            with b2:
                if not cur.empty and st.button(f"🗑️ حذف {label}",key=f"del_{key}",use_container_width=True):
                    try:
                        delete_df(key)
                        log_change("حذف ملف",label)
                        st.success(f"تم حذف ملف {label} نهائيًا من التخزين")
                        st.rerun()
                    except Exception as e:
                        st.error(str(e))

def update_documents():
    st.title("✏️ تحديث وثائق الموظف")
    if not is_hr():
        st.warning("حساب المدير للمشاهدة فقط")
        return

    iq=load_df("iqama")
    ba=load_df("baladi")
    if iq.empty and ba.empty:
        st.info("ارفعي ملف الإقامات أو الكروت البلدية أولًا")
        return

    # Build one employee selector from the uploaded document files.
    people=[]
    seen=set()
    for source,df in [("iqama",iq),("baladi",ba)]:
        if df.empty: continue
        idc=find_col(df,["الرقم الوظيفي","رقم الموظف","employee_id","id"])
        nc=find_col(df,["اسم الموظف","الموظف","الاسم","employee_name","name"])
        for idx,row in df.iterrows():
            emp_id=str(row[idc]).strip() if idc and pd.notna(row[idc]) else ""
            name=str(row[nc]).strip() if nc and pd.notna(row[nc]) else ""
            key=(emp_id or name).strip().lower()
            if not key or key in seen: continue
            seen.add(key); people.append({"key":key,"id":emp_id,"name":name})
    if not people:
        st.info("لم أجد أسماء موظفين في الملفات المرفوعة")
        return

    labels=[f"{x['name']} — {x['id']}" if x['id'] else x['name'] for x in people]
    choice=st.selectbox("اختاري الموظف",range(len(labels)),format_func=lambda i:labels[i])
    person=people[choice]

    def row_for(df):
        if df.empty:return None
        idc=find_col(df,["الرقم الوظيفي","رقم الموظف","employee_id","id"])
        nc=find_col(df,["اسم الموظف","الموظف","الاسم","employee_name","name"])
        if person['id'] and idc:
            m=df[idc].astype(str).str.strip()==person['id']
            if m.any(): return df.index[m][0]
        if person['name'] and nc:
            m=df[nc].astype(str).str.strip().str.lower()==person['name'].lower()
            if m.any(): return df.index[m][0]
        return None

    irow=row_for(iq); brow=row_for(ba)
    iq_end=find_col(iq,["تاريخ انتهاء الإقامة","انتهاء الإقامة","تاريخ انتهاء الاقامة","iqama_end"]) if not iq.empty else None
    iq_h=find_col(iq,["تاريخ انتهاء الإقامة بالهجري","تاريخ انتهاء الاقامة بالهجري","انتهاء الإقامة هجري","iqama_end_hijri"]) if not iq.empty else None
    pp_end=find_col(iq,["تاريخ انتهاء الجواز","انتهاء الجواز","تاريخ نهاية الجواز","passport_end"]) if not iq.empty else None
    ba_end=find_col(ba,["تاريخ انتهاء الكرت البلدي","تاريخ انتهاء الكرت","انتهاء الكرت البلدي","تاريخ انتهاء البطاقة البلدية","تاريخ الانتهاء","baladi_end"]) if not ba.empty else None

    def old(df,row,col):
        if row is None or col is None or pd.isna(df.at[row,col]): return ""
        v=df.at[row,col]
        if isinstance(v,(pd.Timestamp,datetime,date)): return f"{v.year:04d}/{v.month:02d}/{v.day:02d}"
        return str(v).split()[0]

    st.caption("اكتبي التاريخ كما هو في الوثيقة: مثال ميلادي 2027/10/25 أو هجري 1449/04/08. النظام يتعرف عليه تلقائيًا.")
    with st.form("edit_documents_form"):
        c1,c2=st.columns(2)
        iq_val=c1.text_input("🪪 تاريخ انتهاء الإقامة",value=old(iq,irow,iq_end),disabled=irow is None)
        pp_val=c2.text_input("🛂 تاريخ انتهاء الجواز",value=old(iq,irow,pp_end),disabled=irow is None)
        hijri_val=None
        if iq_h is not None:
            hijri_val=c1.text_input("🪪 تاريخ انتهاء الإقامة بالهجري",value=old(iq,irow,iq_h),disabled=irow is None)
        ba_val=c2.text_input("🏥 تاريخ انتهاء الكرت البلدي",value=old(ba,brow,ba_end),disabled=brow is None)
        ok=st.form_submit_button("💾 حفظ تحديث الوثائق",use_container_width=True)

    if ok:
        changed=[]
        if irow is not None:
            if iq_end is None:
                iq_end="تاريخ انتهاء الإقامة"; iq[iq_end]=""
            if pp_end is None:
                pp_end="تاريخ انتهاء الجواز"; iq[pp_end]=""
            iq.at[irow,iq_end]=iq_val.strip()
            iq.at[irow,pp_end]=pp_val.strip()
            if iq_h is not None and hijri_val is not None: iq.at[irow,iq_h]=hijri_val.strip()
            save_df("iqama",iq); changed.append("الإقامة/الجواز")
        if brow is not None:
            if ba_end is None:
                ba_end="تاريخ انتهاء الكرت البلدي"; ba[ba_end]=""
            ba.at[brow,ba_end]=ba_val.strip()
            save_df("baladi",ba); changed.append("الكرت البلدي")
        if changed:
            log_change("تحديث وثائق موظف",f"{person['name']} - " + "، ".join(changed))
            st.success("تم تحديث الوثائق وحفظها. ستتحدث التنبيهات تلقائيًا.")
            st.rerun()
        else:
            st.warning("لم أجد سجلًا لهذا الموظف في ملفات الوثائق")

def system_admin():
    st.title("⚙️ إدارة النظام")
    if not is_hr():st.warning("غير متاح لحساب المدير");return
    st.write("عقود الموظفين: 90 يوم")
    st.write("الإقامات والجوازات: 90 يوم (3 أشهر)")
    st.write("الكروت البلدية: 20 يوم")
    st.write("الإيجارات: 30 يوم قبل الاستحقاق")
    st.info("يمكن إضافة أقسام أو خانات جديدة من هذا القسم في الإصدارات القادمة.")
    st.warning("بيانات الدخول: HR = hr / 1234، Manager = manager / Manager@2026")

if not st.session_state.get("logged"):login();st.stop()

with st.sidebar:
    st.markdown("## 👥 HR System");st.caption(f"{st.session_state.display_name} — {'تعديل كامل' if is_hr() else 'مشاهدة فقط'}")
    pages=["لوحة التحكم","الموظفون","إضافة موظف","نقل موظف وسجل الحركات","أرشيف الموظفين","عقود الموظفين","الإقامات","الكروت البلدية","تحديث الوثائق","الإيجارات","التنبيهات","سجل التعديلات","إدارة البيانات","إدارة النظام"]
    if not is_hr():pages=[p for p in pages if p not in {"إضافة موظف","تحديث الوثائق","إدارة البيانات","إدارة النظام"}]
    page=st.radio("القائمة",pages)
    st.markdown("---")
    if st.button("تسجيل الخروج",use_container_width=True):st.session_state.clear();st.rerun()

if page=="لوحة التحكم":dashboard()
elif page=="الموظفون":data_page("employees","👥 الموظفون")
elif page=="إضافة موظف":add_employee()
elif page=="نقل موظف وسجل الحركات":move_employee()
elif page=="أرشيف الموظفين":archive_employee()
elif page=="عقود الموظفين":data_page("contracts","📄 عقود الموظفين")
elif page=="الإقامات":data_page("iqama","🪪 الإقامات")
elif page=="الكروت البلدية":data_page("baladi","🏪 الكروت البلدية")
elif page=="تحديث الوثائق":update_documents()
elif page=="الإيجارات":rents_page()
elif page=="التنبيهات":dashboard()
elif page=="سجل التعديلات":data_page("changes","📜 سجل التعديلات")
elif page=="إدارة البيانات":uploads()
elif page=="إدارة النظام":system_admin()
